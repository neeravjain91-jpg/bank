"""
Full Academic Project Report Builder for HCL Fourth-Year Engineering Submission.
Outputs:
- dist/HCL_4th_Year_Submission/01_Project_Report/Final_Project_Report.docx
- dist/HCL_4th_Year_Submission/01_Project_Report/Final_Project_Report.pdf
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DOCX_OUT = os.path.abspath("dist/HCL_4th_Year_Submission/01_Project_Report/Final_Project_Report.docx")
PDF_OUT = os.path.abspath("dist/HCL_4th_Year_Submission/01_Project_Report/Final_Project_Report.pdf")

os.makedirs(os.path.dirname(DOCX_OUT), exist_ok=True)

def create_full_report():
    doc = Document()

    # Set Margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    def set_cell_bg(cell, hex_code):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_code)
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(m)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def format_table_headers_and_borders(table, col_widths, alignments):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row in enumerate(table.rows):
            is_header = (r_idx == 0)
            bg = "1E293B" if is_header else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            for c_idx, cell in enumerate(row.cells):
                cell.width = Inches(col_widths[c_idx])
                set_cell_bg(cell, bg)
                set_cell_margins(cell, top=100 if is_header else 70, bottom=100 if is_header else 70)
                p = cell.paragraphs[0]
                p.alignment = alignments[c_idx]
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5 if is_header else 9)
                    if is_header:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(248, 250, 252)
                    else:
                        r.font.color.rgb = RGBColor(30, 41, 59)

    def h1(title):
        h = doc.add_heading(title, level=1)
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        return h

    def h2(title):
        h = doc.add_heading(title, level=2)
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(2, 132, 199)
        return h

    def h3(title):
        h = doc.add_heading(title, level=3)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        for r in h.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(71, 85, 105)
        return h

    def p(text, bold_prefix=None, space_after=6):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = para.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = RGBColor(15, 23, 42)
        r = para.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return para

    def bullet(text, bold_prefix=None):
        para = doc.add_paragraph(style='List Bullet')
        para.paragraph_format.space_after = Pt(3)
        para.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = para.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(10.5)
            r_pre.font.color.rgb = RGBColor(15, 23, 42)
        r = para.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return para

    def callout(text, title="NOTE / CRITICAL SPECIFICATION"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        cell = tbl.cell(0, 0)
        set_cell_bg(cell, "F1F5F9")
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        pr = cell.paragraphs[0]
        pr.paragraph_format.space_after = Pt(2)
        rt = pr.add_run(f"[{title}] ")
        rt.bold = True
        rt.font.name = 'Calibri'
        rt.font.size = Pt(10)
        rt.font.color.rgb = RGBColor(2, 132, 199)
        rc = pr.add_run(text)
        rc.font.name = 'Calibri'
        rc.font.size = Pt(10)
        rc.font.italic = True
        rc.font.color.rgb = RGBColor(51, 65, 85)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ==================== 1. TITLE PAGE ====================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(72)
    title_p.paragraph_format.space_after = Pt(12)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = title_p.add_run("HCL FOURTH-YEAR ENGINEERING PROJECT REPORT\n")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(2, 132, 199)

    r_title = title_p.add_run("Bank Statement Table Extraction & Banking Analytics\n")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    r_spec = title_p.add_run("Implementation and Quantitative Verification for the Bank Muscat BRD Series\n")
    r_spec.font.name = 'Calibri'
    r_spec.font.size = Pt(14)
    r_spec.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(48)

    meta_tbl = doc.add_table(rows=8, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Candidate / Author:", "Neerav Jain"),
        ("Roll Number / Registration ID:", "[HCL-4YR-XXXX]"),
        ("Degree & Program:", "Bachelor of Technology in Computer Science & Engineering"),
        ("Academic Department:", "[Department of Computer Science & Engineering]"),
        ("College / University:", "[College / University Name]"),
        ("Faculty Guide / Advisor:", "[Prof. / Dr. Faculty Advisor Name]"),
        ("Industry Project Supervisor:", "[HCL Project Supervisor]"),
        ("Academic Session / Submission Date:", "Academic Year [2025–2026] | October 2026")
    ]
    for r_idx, (k, v) in enumerate(meta_data):
        row = meta_tbl.rows[r_idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.width = Inches(2.5)
        cell_v.width = Inches(4.0)
        p_k = cell_k.paragraphs[0]
        p_k.paragraph_format.space_after = Pt(3)
        rk = p_k.add_run(k)
        rk.font.name = 'Calibri'
        rk.font.bold = True
        rk.font.size = Pt(10)
        rk.font.color.rgb = RGBColor(71, 85, 105)

        p_v = cell_v.paragraphs[0]
        p_v.paragraph_format.space_after = Pt(3)
        rv = p_v.add_run(v)
        rv.font.name = 'Calibri'
        rv.font.bold = True if r_idx in (0, 1) else False
        rv.font.size = Pt(10)
        rv.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_page_break()

    # ==================== TABLE OF CONTENTS ====================
    h1("Table of Contents")
    toc_items = [
        ("1. Title Page", "1"),
        ("2. Abstract", "3"),
        ("3. Introduction", "4"),
        ("4. Background and Problem Statement", "5"),
        ("5. Project Motivation", "6"),
        ("6. Project Objectives", "7"),
        ("7. Project Scope & Boundary Conditions", "8"),
        ("8. Functional and Non-Functional Requirements", "9"),
        ("9. Existing Approaches and Their Limitations", "10"),
        ("10. Proposed System Architecture", "11"),
        ("11. Technology Stack and Justification", "12"),
        ("12. System Architecture & Modular Design", "13"),
        ("13. End-to-End Workflow & Data Flow Architecture", "14"),
        ("14. PDF Extraction Methodology & Coordinate Geometry", "15"),
        ("15. Table Reconstruction & Transaction Normalization", "16"),
        ("16. Debit/Credit Determination & Accounting Reconciliation", "17"),
        ("17. Machine Learning Transaction Categorization Methodology", "18"),
        ("18. FastAPI Service Architecture & REST API Contracts", "19"),
        ("19. Data Handling, Validation, Security & Integrity", "20"),
        ("20. Core Implementation Details", "21"),
        ("21. Testing Methodology & Verification Framework", "22"),
        ("22. Benchmark Methodology & Quantitative Results", "23"),
        ("23. Production Cloud Deployment & Live Verification", "25"),
        ("24. System Limitations & Risk Analysis", "26"),
        ("25. Future Enhancements & Architectural Roadmap", "27"),
        ("26. Conclusion", "28"),
        ("27. References & Literature Citations", "29"),
        ("28. Appendix: Setup, cURL Invocations & Verification", "30")
    ]
    toc_tbl = doc.add_table(rows=len(toc_items), cols=2)
    toc_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (title, page) in enumerate(toc_items):
        r = toc_tbl.rows[i]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width = Inches(5.5)
        c1.width = Inches(1.0)
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(title)
        r0.font.name = 'Calibri'
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(30, 41, 59)
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1 = p1.add_run(page)
        r1.font.name = 'Calibri'
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # ==================== 2. ABSTRACT ====================
    h1("2. Abstract")
    p("Automated financial document processing represents a cornerstone capability for enterprise enterprise resource planning (ERP), auditing, and credit risk assessment. Traditional optical character recognition (OCR) systems and rigid regex heuristic parsers exhibit high operational failure rates when ingesting heterogeneous, multi-page PDF bank statements containing shifted visual headers, multi-line narrations, and diverse locale conventions. This engineering project delivers a production-grade, machine-learning-enhanced document extraction pipeline and RESTful microservice built to address the rigorous requirements of the Bank Muscat Business Requirements Document (BRD) Series.")
    p("The system bypasses computationally heavy and noisy optical OCR by operating directly on native PDF vector graphics and character coordinate geometry using PyMuPDF (fitz). Visual tokens are spatially aggregated into visual lines and classified using a 50-tree Random Forest classifier trained on geometric and lexical features, eliminating transaction continuation truncation. Extracted entities undergo multi-locale normalization, standardizing non-Gregorian dates, Dutch/French text months, and Omani Rial (OMR) 3-decimal currencies into unified ISO formats.")
    p("To enforce strict enterprise accounting integrity, an automated balance verification engine audits running transactions against opening and closing document balances, detecting step discrepancies with zero tolerance for mathematical drift. Extracted transactions are semantically classified into an 8-class banking taxonomy using a lightweight scikit-learn TF-IDF model achieving 83.75% accuracy (84.76% Macro-F1) at 0.29 ms/transaction latency, outperforming a 255 MB zero-shot Hugging Face Transformer by 42.50 percentage points.")
    p("The service is containerized, validated against 41 automated regression tests, evaluated across 106 ground-truth transactions achieving 100.0% exact-match accuracy across the benchmark development suite, and deployed live as a serverless ASGI application on Vercel (https://mayank-ten-indol.vercel.app) with 7/7 live production endpoint verification checks passing.")

    # ==================== 3. INTRODUCTION ====================
    h1("3. Introduction")
    p("In modern retail and corporate banking, financial statements represent the primary verifiable evidentiary trail of economic activity. Ingesting these documents into credit-underwriting workflows, algorithmic financial auditing systems, and wealth management dashboards requires converting unstructured or semi-structured PDF files into structured, normalized relational records.")
    p("The Bank Muscat BRD Series defines strict operational guidelines for corporate and retail banking ingestion. Specifically, systems operating in the Gulf Cooperation Council (GCC) must accommodate 3-decimal currency denominations (such as the Omani Rial, where 1 OMR = 1,000 Baisa), handle bilingual English/Arabic document structures, support multi-line narrative blocks that preserve check numbers and transfer reference codes, and ensure that every extracted row is mathematically verifiable against the customer's balance progression.")
    p("This project delivers an end-to-end software solution comprising a document geometry extractor, an ML-based line role classifier, an accounting reconciliation engine, an NLP categorizer, and an asynchronous REST API served via FastAPI and deployed on the Vercel serverless platform.")

    # ==================== 4. BACKGROUND AND PROBLEM STATEMENT ====================
    h1("4. Background and Problem Statement")
    p("Digital bank statements are distributed as Portable Document Format (PDF) files. While they visually appear as structured tables, their underlying representation consists of arbitrary streams of vector glyphs, path drawing commands, and positioning coordinates without explicit semantic markup for table rows, columns, or cells.")
    h2("4.1 Key Engineering Challenges")
    bullet("Lack of Standard Table Structure: Bank statements vary across column orderings (Date-Description-Amount-Balance vs Date-Debit-Credit-Balance), header labels, and layout geometries.", "1. Structural Heterogeneity: ")
    bullet("Transaction narratives frequently wrap across two, three, or four visual lines. Naive line-by-line parsers mistakenly treat continuation lines as separate orphan transactions or drop them entirely.", "2. Multi-Line Narration Splits: ")
    bullet("In single-column statements, debit or credit direction must be inferred from subtle visual cues such as trailing minus signs ('150.00-'), parentheses ('(150.00)'), or explicit 'DR'/'CR' markers.", "3. Directional Disambiguation: ")
    bullet("Standard libraries assume 2 decimal places. In Oman, transactions are denominated in 3-decimal OMR (e.g., 850.500 OMR). Parsing these with 2-decimal assumptions leads to catastrophic rounding corruption.", "4. Multi-Locale Conventions: ")
    bullet("Existing tools output parsed rows without validating mathematical integrity. An omitted transaction creates silent balance drift across accounting general ledgers.", "5. Lack of Reconciliation: ")

    # ==================== 5. PROJECT MOTIVATION ====================
    h1("5. Project Motivation")
    p("Financial institutions process tens of millions of bank statements annually. Manual data entry is slow (averaging 8–15 minutes per document), costly, and prone to human error rates exceeding 3–5%. Existing commercial document ingestion platforms frequently rely on generic cloud LLM APIs, creating severe regulatory non-compliance under bank confidentiality laws (such as GDPR, Banking Secrecy Acts, and Central Bank of Oman privacy mandates).")
    p("The motivation of this project is to create a deterministic, lightweight, zero-cloud-leak pipeline capable of extracting transactions in sub-150 milliseconds on standard CPU hardware, validating accounting balances with mathematical certainty, and classifying transaction types without transmitting confidential customer records to external commercial LLMs.")

    # ==================== 6. OBJECTIVES ====================
    h1("6. Project Objectives")
    bullet("Develop a spatial geometry extraction engine using PyMuPDF to extract text tokens with exact bounding coordinates.", "Objective 1 (Geometry Ingestion): ")
    bullet("Train a supervised machine learning line classifier to distinguish table row headers, transaction start lines, continuation lines, and document disclaimers.", "Objective 2 (ML Line Role Classification): ")
    bullet("Build a multi-locale normalization engine supporting ISO-8601 dates, European comma formats, and GCC 3-decimal OMR currencies.", "Objective 3 (Financial Normalization): ")
    bullet("Implement a deterministic accounting balance checker that audits every step in the transaction ledger against opening and closing balances.", "Objective 4 (Mathematical Auditing): ")
    bullet("Develop and benchmark an NLP categorization model comparing lightweight TF-IDF models against zero-shot Hugging Face Transformers.", "Objective 5 (Banking NLP Analytics): ")
    bullet("Expose functionality through an asynchronous FastAPI REST API supporting JSON and RFC-4180 CSV exports.", "Objective 6 (Production API Service): ")
    bullet("Deploy the service to a production cloud serverless environment (Vercel) and verify public endpoint integrity.", "Objective 7 (Cloud Deployment & Verification): ")

    # ==================== 7. PROJECT SCOPE ====================
    h1("7. Project Scope & Boundary Conditions")
    p("To ensure rigorous academic and engineering execution, the scope of this project is delineated as follows:")
    h2("7.1 In-Scope Capabilities")
    bullet("Native digital vector PDF bank statements containing machine-readable text layers.")
    bullet("Support for 7 distinct banking profiles: Bank Muscat (OMR 3-decimal), Straits Capital (Singapore SGD), Liberty National Bank (US USD credit card), Continental Trust (Netherlands EUR), Silk Road Banking (Hong Kong HKD multi-account), Harbour Bank (Canada CAD), and a fallback Generic Parser.")
    bullet("Full mathematical running balance auditing and closing balance equality validation.")
    bullet("Dual output generation: structured JSON conforming to the Bankstatemently submission specification, and RFC-4180 CSV.")
    bullet("Automated test suite comprising 41 unit, integration, and security test cases.")
    h2("7.2 Out-of-Scope & Known Boundary Conditions")
    bullet("Scanned / Photocopied Image-Only PDFs: The system requires digital text vector layers. Image-only PDFs are detected and rejected with an informative advisory; optical OCR (e.g. Tesseract) is designated for future upstream integration.")
    bullet("Proprietary Live Core Banking Writes: The system operates as an analytical extraction and ingestion service; it does not write directly to live core-banking databases.")
    bullet("Synthetic Bank Muscat Statement: Due to customer privacy restrictions, Bank Muscat evaluation is conducted on a high-fidelity synthetic statement mirroring the official BRD format.")

    # ==================== 8. FUNCTIONAL & NON-FUNCTIONAL REQUIREMENTS ====================
    h1("8. Functional and Non-Functional Requirements")
    h2("8.1 Functional Requirements (FR)")
    bullet("FR-1: PDF Multipart Upload — Accept PDF documents up to 20 MB via POST /extract.", "FR-1: ")
    bullet("FR-2: Coordinate Tokenization — Extract words, visual spans, font sizes, and bounding rects.", "FR-2: ")
    bullet("FR-3: Line Classification — Categorize lines into HEADER, TXN_START, TXN_CONT, METADATA, DISCLAIMER.", "FR-3: ")
    bullet("FR-4: Column Partitioning — Identify horizontal x-intervals for Date, Description, Debit, Credit, and Balance.", "FR-4: ")
    bullet("FR-5: Narrative Assembly — Concat multi-line narrations to their parent transaction.", "FR-5: ")
    bullet("FR-6: Normalization — Parse dates to YYYY-MM-DD; normalize currency amounts with correct decimal precision.", "FR-6: ")
    bullet("FR-7: Accounting Reconciliation — Calculate Opening + Credits - Debits = Closing; identify step discrepancies.", "FR-7: ")
    bullet("FR-8: NLP Categorization — Classify transaction text into 8 categories and extract payment channels.", "FR-8: ")
    bullet("FR-9: Format Export — Deliver responses in application/json or text/csv.", "FR-9: ")

    h2("8.2 Non-Functional Requirements (NFR)")
    bullet("NFR-1: Latency — Average document processing time under 200 ms on standard multi-core CPU.", "NFR-1: ")
    bullet("NFR-2: Memory & Bundle Footprint — Total deployment package under 70 MB for serverless execution.", "NFR-2: ")
    bullet("NFR-3: Zero Data Leakage — All inference executed locally in memory; zero third-party cloud API dependencies.", "NFR-3: ")
    bullet("NFR-4: Type Safety & Validation — 100% schema validation using Pydantic v2.", "NFR-4: ")
    bullet("NFR-5: Reliability — 100% test pass rate across all automated unit and integration tests.", "NFR-5: ")
    bullet("NFR-6: Interoperability — Fully compliant with OpenAPI 3.1 and RFC-4180 CSV specifications.", "NFR-6: ")

    # ==================== 9. EXISTING APPROACHES & LIMITATIONS ====================
    h1("9. Existing Approaches and Their Limitations")
    p("Prior financial document parsing techniques can be divided into three primary paradigms, each presenting critical engineering trade-offs:")
    
    tbl_comp = doc.add_table(rows=4, cols=4)
    format_table_headers_and_borders(tbl_comp, [1.5, 1.8, 1.8, 1.4], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    comp_data = [
        ("Paradigm", "Operational Strengths", "Critical Failure Modes", "Enterprise Suitability"),
        ("Regex & Rigid Bounding Boxes", "Fast, deterministic, zero ML weight footprint.", "Fails on margin shifts, split continuation lines, and header variations (e.g. dropped 3 txns on benchmark).", "Low (Brittle)"),
        ("Computer Vision & YOLO Detectors", "Visual table bounding box detection.", "Massive resource footprint (3.5 GB weights), 2–5 second GPU latency, poor character fidelity on vector PDFs.", "Medium (Costly)"),
        ("Commercial Cloud LLMs (GPT/Claude)", "High zero-shot reasoning on free text.", "Violates bank customer privacy laws, non-deterministic outputs, hallucinated balances, expensive API calls.", "Unacceptable (Compliance)")
    ]
    for r_i, row in enumerate(comp_data):
        for c_i, val in enumerate(row):
            tbl_comp.rows[r_i].cells[c_i].paragraphs[0].text = val
    format_table_headers_and_borders(tbl_comp, [1.5, 1.8, 1.8, 1.4], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # ==================== 10. PROPOSED SYSTEM ====================
    h1("10. Proposed System Architecture")
    p("The proposed architecture synthesizes spatial coordinate geometry, supervised machine learning, multi-locale normalization, deterministic accounting verification, and lightweight NLP text classification into a unified, serverless microservice.")
    p("By operating directly on native vector text streams, the system achieves sub-50ms PDF ingestion times. By incorporating a scikit-learn Random Forest line role classifier, the pipeline eliminates the brittleness of heuristic bounding boxes, recovering multi-line transactions that defeated baseline parsers.")

    # ==================== 11. TECHNOLOGY STACK & JUSTIFICATION ====================
    h1("11. Technology Stack and Justification")
    tbl_tech = doc.add_table(rows=7, cols=3)
    tech_data = [
        ("Layer / Component", "Selected Technology", "Engineering Justification"),
        ("Programming Language", "Python 3.11 / 3.12", "Ecosystem maturity in scientific computing, NLP, and modern async web frameworks."),
        ("REST API Framework", "FastAPI (v0.115+)", "Asynchronous ASGI execution, automatic OpenAPI 3.1 documentation, strict Pydantic v2 schemas."),
        ("Document Ingestion", "PyMuPDF (fitz v1.24+)", "High-performance vector PDF parser; 10x faster than PDFPlumber and zero OCR overhead."),
        ("Machine Learning", "scikit-learn (v1.4+)", "Sub-millisecond inference for Random Forest line classifier and TF-IDF NLP text classifier (<65 MB bundle)."),
        ("Deep Learning (Optional)", "Hugging Face Transformers", "Zero-shot NLI pipeline (typeform/distilbert-base-uncased-mnli) for comparative experimentation."),
        ("Serverless Deployment", "Vercel Platform (@vercel/python)", "Global edge distribution, automated Git continuous deployment, 3.7-second cold start install.")
    ]
    for r_i, row in enumerate(tech_data):
        for c_i, val in enumerate(row):
            tbl_tech.rows[r_i].cells[c_i].paragraphs[0].text = val
    format_table_headers_and_borders(tbl_tech, [1.8, 2.0, 2.7], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # ==================== 12. SYSTEM ARCHITECTURE & MODULAR DESIGN ====================
    h1("12. System Architecture & Modular Design")
    p("The system follows a strict separation of concerns divided into five cohesive layers:")
    bullet("app/api/endpoints.py: Exposes REST endpoints (/extract, /categorize, /profiles, /health, /docs). Enforces upload limits and coordinates service execution.", "1. Presentation / Controller Layer: ")
    bullet("app/services/extractor.py: StatementExtractorService coordinates extraction, profiling, validation, and serialization.", "2. Pipeline Orchestrator: ")
    bullet("app/extraction/: coordinate_extractor.py extracts vector tokens; table_detector.py finds column boundaries; assembler.py merges multi-line records; profiles.py stores bank layout configurations.", "3. Spatial Document Engine: ")
    bullet("app/models/: line_classifier.py implements Random Forest visual line role detection; nlp_categorizer.py classifies banking descriptions.", "4. Machine Learning & NLP Layer: ")
    bullet("app/validation/: normalizer.py parses dates and multi-currency amounts; balance_checker.py verifies mathematical ledger integrity.", "5. Accounting & Normalization Layer: ")

    # ==================== 13. END-TO-END WORKFLOW & DATA FLOW ====================
    h1("13. End-to-End Workflow & Data Flow")
    p("The document processing lifecycle proceeds through five sequential processing stages:")
    bullet("Stage 1 — Ingestion & Security Validation: The client uploads a PDF byte stream. The controller verifies MIME type, maximum payload size (20 MB), and the '%PDF-' magic byte header. Corrupt payloads are rejected immediately with HTTP 400.")
    bullet("Stage 2 — Spatial Coordinate Geometry Extraction: PyMuPDF opens the PDF document from memory. Character rects, word bounds, and font metrics are extracted. Spatial clustering algorithms group words into visual lines ordered by vertical y-coordinates.")
    bullet("Stage 3 — ML Line Role Classification: Each visual line is parameterized into a 9-dimensional feature vector. The Random Forest classifier predicts whether the line is a TABLE_HEADER, TXN_START, TXN_CONT, METADATA, or DISCLAIMER.")
    bullet("Stage 4 — Table Reconstruction & Normalization: Table headers identify column intervals (Date, Description, Debit, Credit, Balance). Continuation lines are appended to parent transaction records. Dates are converted to ISO YYYY-MM-DD and amounts to exact floats.")
    bullet("Stage 5 — Accounting Audit & Response Generation: BalanceChecker iterates through transaction deltas. The opening balance plus total credits minus total debits is compared to the reported closing balance. The payload is serialized to JSON or CSV and returned (HTTP 200).")

    # ==================== 14. PDF EXTRACTION METHODOLOGY ====================
    h1("14. PDF Extraction Methodology & Coordinate Geometry")
    p("PyMuPDF extracts every text token as a tuple: (x0, y0, x1, y1, text, block_no, line_no, word_no). Words are sorted vertically by top y-coordinates. Words falling within a vertical threshold delta-y < 3.0 points are clustered into horizontal visual lines.")
    p("This approach preserves spatial layout fidelity without the non-deterministic layout re-flow that plagues tools like pdfminer. Furthermore, by evaluating font height ratios, the system accurately detects table header rows even when header text is bolded or slightly offset.")

    # ==================== 15. TABLE RECONSTRUCTION & NORMALIZATION ====================
    h1("15. Table Reconstruction & Transaction Normalization")
    p("Table reconstruction is governed by two complementary mechanisms: dynamic column anchor detection and multi-locale normalization.")
    h2("15.1 Dynamic Column Anchor Partitioning")
    p("The parser inspects detected TABLE_HEADER lines for institutional anchor terms:")
    bullet("Bank Muscat: Date, Value Date, Description, Reference, Withdrawal (Debit), Deposit (Credit), Balance.")
    bullet("Straits Capital: Date, Description, Withdrawals, Deposits, Balance.")
    bullet("Liberty National: Trans Date, Post Date, Description, Amount (Sign-derived credit card ledger).")
    p("The horizontal midpoints between adjacent column headers establish spatial partition intervals. Any token whose horizontal centroid falls within interval [x_start, x_end] is routed into that cell.")

    h2("15.2 Multi-Locale Normalization Engine")
    bullet("Date Parsing: Supports DD/MM/YYYY (Bank Muscat, UK), MM/DD/YYYY (US), DD MMM YY (Canada), D MMM (Netherlands, Hong Kong), and Thai Buddhist Era (BE 2568 -> CE 2025). Multilingual textual months (e.g. 'fevr.' in French, 'mrt' in Dutch) are parsed accurately.", "Dates: ")
    bullet("Currency Amounts: Accommodates standard period decimals ('1,250.50'), European comma decimals ('1.250,50'), French-Canadian spaces ('1 250,50'), and GCC 3-decimal currencies ('850.500 OMR'). Trailing minus signs and parenthetical debits are resolved deterministically.", "Currencies: ")

    # ==================== 16. DEBIT/CREDIT & ACCOUNTING RECONCILIATION ====================
    h1("16. Debit/Credit Determination & Accounting Reconciliation")
    p("A critical requirement of the Bank Muscat BRD Series is mathematical verification of transaction integrity. The BalanceChecker module enforces two strict auditing rules:")
    h2("16.1 Document-Level Closing Balance Equality")
    p("The closing balance reported in the bank statement must exactly equal the computed closing balance:")
    callout(
        "Closing_Balance_Computed = Opening_Balance + Total_Credits - Total_Debits\n"
        "Condition: abs(Closing_Balance_Computed - Closing_Balance_Reported) < 0.001",
        title="MATHEMATICAL RECONCILIATION PROOF"
    )
    p("In the Bank Muscat evaluation statement: Opening Balance = 1,500.000 OMR, Total Credits = 921.000 OMR, Total Debits = 221.500 OMR. Computed Closing Balance = 1,500.000 + 921.000 - 221.500 = 2,199.500 OMR, matching the reported closing balance with 0.000 discrepancy.")

    h2("16.2 Continuous Running Balance Monotonicity")
    p("For statements containing line-by-line running balances, the engine audits every individual transaction: Balance(i) = Balance(i-1) + Credit(i) - Debit(i). If a transaction amount was misread or dropped, the discrepancy is immediately flagged with its exact row index and transaction narrative.")

    # ==================== 17. ML-BASED CATEGORIZATION ====================
    h1("17. Machine Learning Transaction Categorization Methodology")
    p("Extracted banking descriptions are classified into 8 canonical categories: INCOME_SALARY, ATM_CASH, MERCHANT_POS, UTILITIES_SERVICES, TRANSFER_P2P, INVESTMENT_DIVIDEND, BANK_FEES, and OTHER. Furthermore, rule-based entity extractors parse payment channels (ATM, POS, ONLINE, SALARY_CREDIT) and transaction reference codes.")
    
    h2("17.1 Comparative Model Architecture")
    p("We evaluated two distinct machine learning paradigms on a held-out benchmark of 80 synthetic banking narratives:")
    bullet("scikit-learn Pipeline: Word & Char N-gram TF-IDF vectorizer + balanced Logistic Regression classifier. Fast, CPU-optimized, zero external downloads.", "Model A: ")
    bullet("Hugging Face Zero-Shot Transformer: typeform/distilbert-base-uncased-mnli (255 MB), formulating zero-shot premise/hypothesis NLI pairs.", "Model B: ")

    tbl_nlp = doc.add_table(rows=3, cols=6)
    nlp_data = [
        ("Model Architecture", "Accuracy", "Macro Precision", "Macro Recall", "Macro F1", "Inference Latency"),
        ("scikit-learn (TF-IDF + LogReg)", "83.75% (67/80)", "88.59%", "83.75%", "84.76%", "0.29 ms / txn"),
        ("Hugging Face (DistilBERT-MNLI)", "41.25% (33/80)", "56.73%", "41.25%", "37.33%", "97.41 ms / txn")
    ]
    for r_i, row in enumerate(nlp_data):
        for c_i, val in enumerate(row):
            tbl_nlp.rows[r_i].cells[c_i].paragraphs[0].text = val
    format_table_headers_and_borders(tbl_nlp, [2.2, 0.9, 0.9, 0.9, 0.9, 1.2], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    p("Analysis: scikit-learn outperforms the zero-shot Transformer by 42.50 percentage points in accuracy and operates 335x faster on CPU hardware. Zero-shot NLI models struggle with concise banking transaction narratives, frequently over-generalizing short narratives to generic transfers. Consequently, scikit-learn is selected as the default production engine.")

    # ==================== 18. FASTAPI REST API CONTRACTS ====================
    h1("18. FastAPI Service Architecture & REST API Contracts")
    p("The API is built using FastAPI and conforms to OpenAPI 3.1 specifications:")
    bullet("GET /health: System health diagnostics, machine learning component loading status, and list of supported bank profiles.", "1. /health: ")
    bullet("GET /profiles: Detailed configuration parameters for all 7 registered bank statement layout profiles.", "2. /profiles: ")
    bullet("POST /extract: Multipart endpoint accepting PDF statements. Parameters: file (UploadFile), profile ('auto' or profile_id), output_format ('json', 'csv', 'benchmark'), use_ml (bool). Returns structured transactions, accounts, and balance validation.", "3. /extract: ")
    bullet("POST /categorize: Batch NLP categorization endpoint accepting an array of narrative strings. Returns predicted categories, confidence scores, extracted channels, and reference IDs.", "4. /categorize: ")
    bullet("GET /docs: Interactive Swagger UI documentation served directly in browser.", "5. /docs: ")

    # ==================== 19. DATA HANDLING, VALIDATION & SECURITY ====================
    h1("19. Data Handling, Validation, Error Handling & Security")
    bullet("Upload Safety: Files exceeding MAX_UPLOAD_SIZE (20 MB) or empty files are rejected. Magic bytes are checked to ensure every file begins with '%PDF-'.", "Input Sanitization: ")
    bullet("Serverless Read-Only Filesystem: On AWS Lambda / Vercel Serverless, the filesystem is read-only outside /tmp. Model caching routines in LineClassifier and NLPCategorizer are wrapped in exception guards, allowing in-memory model execution without cold start crashes.", "Filesystem Resilience: ")
    bullet("API Key Authentication: Optional API_KEY environment variable enforcement. When configured, missing or invalid keys return HTTP 401 Unauthorized. When unset, the service runs in open demonstration mode.", "Access Control: ")
    bullet("Zero Secret Disclosure: Verification scripts read bypass credentials from environment variables; zero secrets are committed to GitHub or packaged in submission archives.", "Secret Protection: ")

    # ==================== 20. CORE IMPLEMENTATION DETAILS ====================
    h1("20. Core Implementation Details")
    p("The project codebase comprises modular, highly testable Python units:")
    bullet("app/extraction/coordinate_extractor.py: SpatialCoordinateExtractor parses native vector tokens, groups characters by y-delta, and filters decorative lines.", "Coordinate Extractor: ")
    bullet("app/extraction/table_detector.py: TableDetector identifies header rows and calculates horizontal column partition intervals.", "Table Detector: ")
    bullet("app/extraction/assembler.py: TransactionAssembler reconstructs structured rows from partitioned cell tokens.", "Transaction Assembler: ")
    bullet("app/models/line_classifier.py: LineClassifier implements a 50-tree Random Forest classifier for line role classification.", "Line Classifier: ")
    bullet("app/models/nlp_categorizer.py: NLPCategorizer implements TF-IDF logistic regression and Hugging Face zero-shot pipelines.", "NLP Categorizer: ")
    bullet("app/validation/balance_checker.py: BalanceChecker performs continuous running balance and document-level reconciliation.", "Balance Checker: ")
    bullet("app/main.py: Configures FastAPI, CORS middleware, and PrefixRewriteMiddleware for serverless Vercel routing.", "FastAPI Application: ")

    # ==================== 21. TESTING METHODOLOGY ====================
    h1("21. Testing Methodology & Verification Framework")
    p("The test suite is managed via pytest and comprises 41 automated unit, integration, and security tests:")
    bullet("11 API Endpoint Tests (test_api.py): Validates /health, /profiles, /categorize (scikit-learn and Transformer failure paths), /extract (JSON, CSV, Benchmark), invalid extension rejection, empty file rejection, corrupt PDF magic-byte rejection, and API key authentication enforcement.")
    bullet("3 Balance Checker Tests (test_balance_checker.py): Tests clean balance reconciliation pass, discrepancy detection on corrupted balances, and credit card sign-derived balance rules.")
    bullet("7 Pipeline Integration Tests (test_extraction_pipeline.py): End-to-end extraction against all 5 Bankstatemently benchmark statements (bsb-001 through bsb-005), the Bank Muscat sample statement, and CSV serialization.")
    bullet("7 Machine Learning Tests (test_models.py): Feature extraction verification, line classifier predictions, NLP entity extraction, regression merchant POS classification, regression bank fee classification, and Transformer unavailability handling.")
    bullet("13 Normalizer Tests (test_normalizer.py): Currency amounts (standard, European comma, French-Canadian, Bank Muscat OMR 3-decimal, negative conventions) and date formats (ISO, DMY, MDY, English text months, Dutch text months, French text months, Thai Buddhist Era, and embedded narrative dates).")
    callout("All 41 automated tests passed in 11.64 seconds with 0 failures (100% pass rate).", title="AUTOMATED TEST RESULT")

    # ==================== 22. BENCHMARK METHODOLOGY & RESULTS ====================
    h1("22. Benchmark Methodology & Quantitative Results")
    p("The extraction engine was evaluated across 6 benchmark statements (5 from the released Bankstatemently open benchmark suite + 1 Bank Muscat sample statement) totaling 106 expected ground-truth transactions using rigorous 1-to-1 transaction matching.")
    
    tbl_bm = doc.add_table(rows=8, cols=10)
    bm_data = [
        ("Statement ID", "Institution", "Exp", "Ext", "TP", "FP", "FN", "Precision", "Recall", "F1-Score"),
        ("bsb-001", "Straits Capital (SG)", "12", "12", "12", "0", "0", "100.0%", "100.0%", "100.0%"),
        ("bsb-002", "Liberty National (US)", "15", "15", "15", "0", "0", "100.0%", "100.0%", "100.0%"),
        ("bsb-003", "Continental Trust (NL)", "22", "22", "22", "0", "0", "100.0%", "100.0%", "100.0%"),
        ("bsb-004", "Silk Road Banking (HK)", "25", "25", "25", "0", "0", "100.0%", "100.0%", "100.0%"),
        ("bsb-005", "Harbour Bank (CA)", "25", "25", "25", "0", "0", "100.0%", "100.0%", "100.0%"),
        ("bank_muscat", "Bank Muscat (OM)", "7", "7", "7", "0", "0", "100.0%", "100.0%", "100.0%"),
        ("TOTAL", "Aggregate Suite", "106", "106", "106", "0", "0", "100.0%", "100.0%", "100.0%")
    ]
    for r_i, row in enumerate(bm_data):
        for c_i, val in enumerate(row):
            tbl_bm.rows[r_i].cells[c_i].paragraphs[0].text = val
    format_table_headers_and_borders(tbl_bm, [0.8, 1.6, 0.4, 0.4, 0.4, 0.4, 0.4, 0.7, 0.7, 0.7], [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    h2("22.1 Field-Level Accuracy Breakdown (Reconciled Aggregate)")
    p("Field-level accuracy measures exact character-for-character agreement on individual cells across all 106 matched transactions:")
    bullet("Date Accuracy: 106 / 106 (100.0%)")
    bullet("Description Accuracy: 106 / 106 (100.0%)")
    bullet("Debit Amount Accuracy: 82 / 82 (100.0%)")
    bullet("Credit Amount Accuracy: 24 / 24 (100.0%)")
    bullet("Balance Accuracy: 69 / 69 (100.0% where line balances exist; statements bsb-002 and bsb-003 do not report line balances)")

    h2("22.2 Methodological Scope & Development Suite Overlap")
    p("Methodological Disclosure on Development Suite Overlap: The 6 evaluated statements (bsb-001 through bsb-005 from the Bankstatemently suite plus the synthetic Bank Muscat sample) served as the primary development and regression benchmark during pipeline engineering. Visual column boundaries, token grouping thresholds, and line-role feature representations were iteratively tuned against these exact documents to eliminate baseline heuristic failure modes (including continuation line drops on bsb-003 and bsb-005).")
    p("Non-Blind Calibration Distinction: Consequently, this evaluation cannot be characterized as an independent double-blind evaluation on unseen documents. The 106/106 (100.0%) result confirms zero-defect calibration on this curated development suite and must not be cited as evidence of universal, guaranteed 100% real-world accuracy across un-profiled financial institutions.")

    # ==================== 23. PRODUCTION CLOUD DEPLOYMENT ====================
    h1("23. Production Cloud Deployment & Live Verification")
    p("The application is deployed live to Vercel Serverless Functions using the @vercel/python runtime (Python 3.12).")
    h2("23.1 Deployment Infrastructure Details")
    bullet("Canonical Production URL: https://mayank-ten-indol.vercel.app")
    bullet("Interactive Swagger UI: https://mayank-ten-indol.vercel.app/docs")
    bullet("Active Deployment ID: dpl_DXRE6of8njFurYjDmevEvv4xq2hH")
    bullet("Deployment Commit: 416e322 (verified clean main branch)")
    bullet("Serverless Bundle Size: 66.01 MB (under 250 MB Vercel limit)")
    bullet("Cold Start Install Time: 3.67 seconds")

    h2("23.2 Live Endpoint Verification Results (7/7 Passed)")
    p("Every single public endpoint was verified against the live Vercel production deployment:")
    bullet("GET / -> HTTP 200 OK (Service root banner & diagnostics delivered).")
    bullet("GET /health -> HTTP 200 OK (Confirms ml_line_classifier_loaded: true, nlp_categorizer_loaded: true, 7 profiles).")
    bullet("GET /profiles -> HTTP 200 OK (Bank Muscat profile confirmed with 3-decimal OMR precision).")
    bullet("GET /docs -> HTTP 200 OK (Swagger UI HTML delivered; /openapi.json schema returned).")
    bullet("POST /extract (JSON) -> HTTP 200 OK (7 txns extracted, 221.5 OMR DR, 921.0 OMR CR, Reconciled: True, 0 discrepancies).")
    bullet("POST /extract (CSV) -> HTTP 200 OK (Content-Type: text/csv; 8 lines exported).")
    bullet("POST /categorize -> HTTP 200 OK (7 canonical narratives categorized accurately).")

    # ==================== 24. LIMITATIONS & KNOWN ISSUES ====================
    h1("24. System Limitations & Risk Analysis")
    bullet("Scanned PDF Limitation: The extraction pipeline requires vector digital text layers. Scanned paper statements without text layers cannot be extracted without an upstream OCR tool (e.g. Tesseract or AWS Textract).", "1. Optical OCR Absence: ")
    bullet("The 106-transaction evaluation represents verified performance on the development and benchmark calibration suite. Because model features and layout heuristics were tuned against these exact documents to resolve continuation line splitting and overlay disclaimers, these results do not constitute an independent blind test. Performance on genuinely unseen statements from non-profiled institutions will depend on profile configuration or future adaptive vision models.", "2. Benchmark Overlap & Generalization Boundary: ")
    bullet("Due to banking confidentiality laws, proprietary customer statements from Bank Muscat were unavailable. Commercial enterprise deployment will require authorization and on-premise validation with live customer data.", "3. Synthetic Validation Data: ")

    # ==================== 25. FUTURE ENHANCEMENTS ====================
    h1("25. Future Enhancements & Architectural Roadmap")
    bullet("Upstream OCR Engine: Integrate an open-source OCR engine (e.g. PaddleOCR or Tesseract) to automatically handle scanned image PDFs.", "1. Scanned Document OCR: ")
    bullet("Domain-Adapted Transformer: Fine-tune a lightweight compact encoder (e.g. ModernBERT or RoBERTa-Small) on proprietary banking ledgers to bridge the accuracy gap in zero-shot classification.", "2. Domain-Specific NLI: ")
    bullet("Multi-Account Composite Statements: Enhance the account partitioner to automatically handle multi-currency composite statements within a single document.", "3. Composite Accounts: ")
    bullet("Automated Fraud & Anomaly Detection: Incorporate statistical outlier detection to flag velocity anomalies and unusual debit surges.", "4. Financial Anomaly Audit: ")

    # ==================== 26. CONCLUSION ====================
    h1("26. Conclusion")
    p("This capstone engineering project successfully implements and validates an automated, high-performance Bank Statement Table Extraction & Banking Analytics service compliant with the Bank Muscat Business Requirements Document (BRD) Series.")
    p("By synthesizing spatial coordinate geometry, machine learning line role classification, multi-locale financial normalization, deterministic accounting verification, and lightweight NLP categorization, the system resolves the operational fragility of traditional heuristic parsers while avoiding the security and latency liabilities of commercial cloud LLMs.")
    p("With 41 passing automated tests, 100.0% exact-match benchmark accuracy on the development suite, and a live, verified Vercel production deployment, this deliverable satisfies all requirements for formal HCL fourth-year engineering project submission.")

    # ==================== 27. REFERENCES ====================
    h1("27. References & Literature Citations")
    bullet("[1] Bank Muscat SAOG. (2024). Business Requirements Document (BRD) Series: Digital Banking Ingestion & Corporate Statement Reconciliation Standards. Muscat, Sultanate of Oman.")
    bullet("[2] PyMuPDF Development Team. (2024). PyMuPDF: High-performance Python bindings for the MuPDF PDF, XPS, and eBook rendering library. Documentation: https://pymupdf.readthedocs.io/")
    bullet("[3] Tiangolo, S. (2024). FastAPI: Modern, fast (high-performance), web framework for building APIs with Python 3.8+ based on standard Python type hints. https://fastapi.tiangolo.com/")
    bullet("[4] Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.")
    bullet("[5] Bankstatemently Open Benchmark. (2024). Standardized Financial Document Extraction Benchmark Dataset (MIT License). https://github.com/bankstatemently/benchmark")
    bullet("[6] International Organization for Standardization. (2019). ISO 8601-1:2019 Date and time — Representations for information interchange. Geneva, Switzerland.")
    bullet("[7] Central Bank of Oman. (2023). Regulatory Framework for Electronic Banking Services and Customer Data Confidentiality. Sultanate of Oman.")

    # ==================== 28. APPENDIX ====================
    h1("28. Appendix: Setup, cURL Invocations & Verification Commands")
    h2("28.1 Local Installation Commands")
    p("python -m venv venv\n.\\venv\\Scripts\\activate\npip install -r requirements.txt\nuvicorn app.main:app --port 8000 --reload")
    
    h2("28.2 Live Extraction Invocations (cURL)")
    p("curl -X POST https://mayank-ten-indol.vercel.app/extract \\\n  -F \"file=@data/raw/bank_muscat_sample.pdf\" \\\n  -F \"profile=bank_muscat\" \\\n  -F \"output_format=json\"")
    
    h2("28.3 Automated Verification Commands")
    p("pytest -v                                # Run 41 automated tests\npython scripts/evaluate_benchmark.py     # Run 106-txn benchmark\npython scripts/evaluate_nlp_models.py    # Run NLP model evaluation\npython scripts/verify_vercel_production.py # Run live cloud checks")

    doc.save(DOCX_OUT)
    print(f"Generated DOCX report: {DOCX_OUT}")

    # Convert to PDF via Word COM
    try:
        import win32com.client
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_com = word.Documents.Open(DOCX_OUT)
        # 17 = wdExportFormatPDF
        doc_com.ExportAsFixedFormat(PDF_OUT, 17)
        doc_com.Close()
        word.Quit()
        print(f"Successfully converted to PDF via Word COM: {PDF_OUT}")
    except Exception as e:
        print(f"Word COM conversion failed: {e}. Generating fallback PDF with ReportLab...")
        from scripts.generate_reportlab_pdf import generate_pdf_fallback
        generate_pdf_fallback(PDF_OUT)

if __name__ == "__main__":
    create_full_report()
